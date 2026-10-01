import json
from copy import deepcopy
from openpyxl import load_workbook

MONTH_MAP = {
    "Enero": "01",
    "Febrero": "02",
    "Marzo": "03",
    "Abril": "04",
    "Mayo": "05",
    "Junio": "06",
    "Julio": "07",
    "Agosto": "08",
    "Septiembre": "09",
    "Octubre": "10",
    "Noviembre": "11",
    "Diciembre": "12",
}


def month_to_date(mes: str) -> str:
    month = MONTH_MAP.get(mes, "01")
    return f"2026-{month}-20T10:00:00.000Z"


def estado(enviada, pagado) -> str:
    if pagado:
        return "cobrado"
    if enviada:
        return "facturado"
    return "pendiente"


def round2(value) -> float:
    if value is None:
        return 0.0
    return round(float(value), 2)


def cliente_concepto(cliente, comentarios) -> tuple[str, str]:
    cliente = (cliente or "").strip()
    comentarios = (comentarios or "").strip()
    # The July ALVAFARM row used comments as concepto in an older backup;
    # keep Cliente as concepto and comments as notas.
    return cliente, comentarios


def main():
    wb = load_workbook("Registro_Autonomo_2026.xlsx", data_only=True)
    ws = wb["Registro_Horas"]

    ingresos = []
    excel_rows = []
    current_id = 1
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not any(row):
            continue
        mes, cliente, horas, precio, total, enviada, pagado, comentarios = row
        if not mes or not cliente:
            continue
        concepto, notas = cliente_concepto(cliente, comentarios)
        ingreso = {
            "id": current_id,
            "concepto": concepto,
            "tipo": "variable",
            "importe": round2(total),
            "fecha": month_to_date(mes),
            "recurrente": False,
            "frecuencia": "unico",
            "notas": notas,
            "estado": estado(enviada, pagado),
        }
        ingresos.append(ingreso)
        excel_rows.append(
            {
                "Mes": mes,
                "Cliente": concepto,
                "Horas": horas,
                "Precio/Hora": precio,
                "Total": round2(total),
                "Enviada": bool(enviada),
                "Pagado": bool(pagado),
                "Comentarios": notas,
            }
        )
        current_id += 1

    with open("analisis_datos/finanzas-backup-2026-07-31.json", encoding="utf-8") as f:
        backup = json.load(f)

    new_backup = deepcopy(backup)
    new_backup["ingresos"] = ingresos

    out_path = "analisis_datos/finanzas-backup-2026-10-01.json"
    with open(out_path, "w", encoding="utf-8") as f:
        json.dump(new_backup, f, ensure_ascii=False, indent=2)

    sheets = {"Registro_Horas": excel_rows}
    with open("excel_data.json", "w", encoding="utf-8") as f:
        json.dump(sheets, f, ensure_ascii=False, indent=2)

    from collections import Counter

    by_month = Counter(i["fecha"][5:7] for i in ingresos)
    by_estado = Counter(i["estado"] for i in ingresos)
    total = sum(i["importe"] for i in ingresos)
    cobrado = sum(i["importe"] for i in ingresos if i["estado"] == "cobrado")
    pendiente = sum(i["importe"] for i in ingresos if i["estado"] != "cobrado")
    print(f"Wrote {out_path}")
    print(f"ingresos={len(ingresos)} total={total:.2f} cobrado={cobrado:.2f} pendiente={pendiente:.2f}")
    print("estados", dict(by_estado))
    print("months", dict(sorted(by_month.items())))


if __name__ == "__main__":
    main()
