from openpyxl import load_workbook

wb = load_workbook("Registro_Autonomo_2026.xlsx", data_only=True)
ws = wb["Registro_Horas"]
rows = []
for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
    if not any(row):
        continue
    mes, cliente, horas, precio, total, enviada, pagado, comentarios = row
    rows.append(
        {
            "row": i,
            "Mes": mes,
            "Cliente": cliente,
            "Horas": horas,
            "Precio": precio,
            "Total": total,
            "Enviada": enviada,
            "Pagado": pagado,
            "Comentarios": comentarios,
        }
    )

print("COUNT", len(rows))
from collections import Counter

print("BY MONTH", dict(Counter(r["Mes"] for r in rows)))
print("--- all ---")
for r in rows:
    print(
        f"{r['Mes']}|{r['Cliente']}|{r['Total']}|pagado={r['Pagado']}|env={r['Enviada']}|{r['Comentarios']}"
    )

print("\n==== Resumen_Mensual ====")
ws2 = wb["Resumen_Mensual"]
for row in ws2.iter_rows(values_only=True):
    print(row)
